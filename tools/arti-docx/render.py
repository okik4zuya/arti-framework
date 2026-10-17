#!/usr/bin/env python3
"""
arti-docx - Markdown -> styled .docx (journal manuscript layout)

Rebuilds the entire document body fresh from a Markdown source on each run,
using a *style template* purely for styles/section-properties/footer (A4,
margins, continuous line numbering, page-number footer, named paragraph/table
styles). Markdown is the source of truth; the .docx is a snapshot rendered
from it, the same way arti-convert pdf treats its HTML/PDF output - never a merge
into an existing .docx's content.

Requires python-docx (see ~/.arti/tools/requirements.txt).

This is intentionally not a full CommonMark implementation - just enough
structure for a manuscript draft: headings, single-line paragraphs with
bold/italic, blockquotes, pipe tables, images, bold Table/Fig captions, code
fences (rendered verbatim as a monospace fallback for an ASCII placeholder
not yet swapped for a real image), and a References section.

Explicit non-goals for v1: nested/numbered lists, footnotes, merging into an
existing docx section-by-section. Extend the converter for these if a
document needs them, rather than reaching for pandoc.

Usage:
    python render.py --md manuscript.md --out manuscript.docx [--template style-template.docx]
"""
import argparse
import os
import re
import sys

import docx
import zotero_fields
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.text.run import Run

DEFAULT_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "default-academic-template.docx")
ZOTERO_TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "default-academic-template-zotero.docx")

# Set by render() only for `--cite zotero`; None keeps the plain-text behaviour byte-for-byte.
# Module-level because add_runs() is called from tables/captions/body alike.
_CITER = None

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
TRAILING_ASIDE_RE = re.compile(r"\s*\*\[[^\]]*\]\*\s*$")
CAPTION_RE = re.compile(r"^\*\*(Table|Fig|Figure)\b")
IMAGE_RE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")


def parse_inline_runs(text):
    """Split text into (text, bold, italic) run tuples on non-nested **bold**/*italic*."""
    runs = []
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False, False))
        token = m.group(0)
        if token.startswith("**"):
            runs.append((token[2:-2], True, False))
        else:
            runs.append((token[1:-1], False, True))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], False, False))
    if not runs:
        runs.append((text, False, False))
    return runs


def add_runs(paragraph, text, force_italic=False, mono=False):
    for run_text, bold, italic in parse_inline_runs(text):
        if not run_text:
            continue
        # With --cite zotero, [LIT: key] markers inside a run become Zotero fields; the
        # surrounding text keeps the run's bold/italic.
        pieces = _split_cites(run_text) if _CITER else [(run_text, None)]
        for piece, keys in pieces:
            if keys:
                for r in _CITER.citation_runs(keys):
                    paragraph._p.append(r)
                continue
            if not piece:
                continue
            run = paragraph.add_run(piece)
            run.bold = bold
            run.italic = italic or force_italic
            if mono:
                run.font.name = "Consolas"
                run.font.size = Pt(9)


def _split_cites(text):
    """[(text, None) | ("", [keys])] -- plain stretches and citation markers, in order."""
    out = []
    pos = 0
    for m in zotero_fields.CITE_RE.finditer(text):
        out.append((text[pos:m.start()], None))
        out.append(("", zotero_fields.split_keys(m.group(1))))
        pos = m.end()
    out.append((text[pos:], None))
    return out


def style_names(document):
    return {s.name for s in document.styles}


def heading_style_for(level, names):
    if level <= 2:
        return "Heading 1"
    if level == 3 and "Heading 3" in names:
        return "Heading 3"
    return "Heading 2"


def table_style_name(document):
    """A house table style from the template, or None. "Table Grid" (all borders) and "Normal Table" do not count:
    without a house style every table gets the three-line academic look (see apply_three_line)."""
    from docx.enum.style import WD_STYLE_TYPE
    for s in document.styles:
        if s.type == WD_STYLE_TYPE.TABLE and s.name not in ("Normal Table", "Table Grid"):
            return s.name
    return None


def text_width(document):
    section = document.sections[-1]
    return section.page_width - section.left_margin - section.right_margin


def page_widths(document):
    """(portrait, landscape) text widths in points, from the last section's margins."""
    section = document.sections[-1]
    margins = section.left_margin + section.right_margin
    short, long_ = sorted((section.page_width, section.page_height))
    return (short - margins) / 12700.0, (long_ - margins) / 12700.0


def is_table_separator(line):
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c) for c in cells)


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def parse_table(lines, i):
    """(header, rows, next_index) for the pipe table whose header line is lines[i]."""
    header = split_row(lines[i])
    rows = []
    j = i + 2
    while j < len(lines) and lines[j].strip().startswith("|"):
        rows.append(split_row(lines[j]))
        j += 1
    return header, rows, j


# Table layout. Widths are estimated in points from character counts (Arial averages about 0.55 em per
# character), because python-docx tables are not auto-fitted when the file is opened.
TABLE_FONT_SIZES = (10, 9, 8)
TABLE_MIN_FONT = 7
CELL_PAD_PT = 10.8          # default left + right cell margin (108 twips each)
CHAR_EM = 0.55
IDEAL_CHARS_CAP = 40        # a long text cell wraps instead of widening its column without limit


def _column_chars(header, rows):
    ideal, minimum = [], []
    for k in range(len(header)):
        cells = [header[k]] + [r[k] for r in rows if k < len(r)]
        plain = [c.replace("*", "") for c in cells]
        ideal.append(min(IDEAL_CHARS_CAP, max((len(part) for c in plain for part in c.split("\n")), default=1)))
        minimum.append(max((len(w) for c in plain for w in re.split(r"[\s/]+", c)), default=1))
    return ideal, minimum


def _widths_pt(header, rows, size):
    ideal, minimum = _column_chars(header, rows)
    em = CHAR_EM * size
    return ([max(c, 3) * em + CELL_PAD_PT for c in ideal], [max(c, 3) * em * 1.25 + CELL_PAD_PT for c in minimum])


def plan_table(header, rows, portrait_pt, landscape_pt):
    """(landscape?, font size). Portrait at 10, 9, 8 pt first; a table that still does not fit goes to a
    landscape section; as a last resort 7 pt in landscape."""
    for landscape, width in ((False, portrait_pt), (True, landscape_pt)):
        for size in TABLE_FONT_SIZES:
            ideal, minimum = _widths_pt(header, rows, size)
            if sum(minimum) <= width and sum(ideal) <= 1.25 * width:
                return landscape, size
    return True, TABLE_MIN_FONT


def allocate_widths(header, rows, size, total_pt):
    ideal, minimum = _widths_pt(header, rows, size)
    if sum(ideal) <= total_pt:
        extra = total_pt - sum(ideal)
        return [w + extra * w / sum(ideal) for w in ideal]
    if sum(minimum) >= total_pt:
        return [w * total_pt / sum(minimum) for w in minimum]
    gap = [i - m for i, m in zip(ideal, minimum)]
    spare = total_pt - sum(minimum)
    return [m + spare * g / sum(gap) for m, g in zip(minimum, gap)]


def _border(tag, val, sz=None):
    el = OxmlElement("w:" + tag)
    el.set(qn("w:val"), val)
    if sz:
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
    return el


def apply_three_line(table):
    """Top and bottom rule on the table and one rule under the header row; no vertical or inner lines."""
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for tag, val in (("top", "single"), ("left", "nil"), ("bottom", "single"), ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")):
        borders.append(_border(tag, val, 8 if val == "single" else None))
    # schema order: ... tblW, jc, tblCellSpacing, tblInd, tblBorders, shd, tblLayout, tblCellMar, tblLook
    anchor = next((tbl_pr.find(qn(t)) for t in ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook") if tbl_pr.find(qn(t)) is not None), None)
    if anchor is not None:
        anchor.addprevious(borders)
    else:
        tbl_pr.append(borders)
    for cell in table.rows[0].cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_borders = OxmlElement("w:tcBorders")
        tc_borders.append(_border("bottom", "single", 4))
        tc_pr.append(tc_borders)


def add_table(document, header, rows, style_name, size, total_pt):
    table = document.add_table(rows=1 + len(rows), cols=len(header))
    if style_name:
        table.style = style_name
    for r, cells in enumerate([header] + rows):
        for k in range(len(header)):
            cell = table.rows[r].cells[k]
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt = para.paragraph_format
            fmt.space_before = Pt(2)
            fmt.space_after = Pt(2)
            fmt.line_spacing = 1.0
            add_runs(para, cells[k] if k < len(cells) else "")
            for el in para._p.findall(qn("w:r")):
                run = Run(el, para)
                run.font.size = Pt(size)
                if r == 0:
                    run.bold = True
    widths = allocate_widths(header, rows, size, total_pt)
    table.autofit = False
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    if tbl_w is not None:
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(int(total_pt * 20)))
    for col, grid in zip(widths, table._tbl.tblGrid.findall(qn("w:gridCol"))):
        grid.set(qn("w:w"), str(int(col * 20)))
    for row in table.rows:
        for k, col in enumerate(widths):
            row.cells[k].width = Pt(col)
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    if not style_name:
        apply_three_line(table)
    return table


def open_section(document, landscape):
    """Start a new-page section, landscape or portrait. Header/footer stay linked; page numbering continues."""
    section = document.add_section(WD_SECTION.NEW_PAGE)
    pg_num = section._sectPr.find(qn("w:pgNumType"))
    if pg_num is not None:
        section._sectPr.remove(pg_num)
    w, h = section.page_width, section.page_height
    if (landscape and w < h) or (not landscape and w > h):
        section.page_width, section.page_height = h, w
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT


def add_image(document, alt_text, path, base_dir, width):
    img_path = path if os.path.isabs(path) else os.path.join(base_dir, path)
    if not os.path.exists(img_path):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, "[missing image: %s]" % path, force_italic=True)
        return
    document.add_picture(img_path, width=width)
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.paragraphs[-1].paragraph_format.keep_with_next = True   # the figure stays with the caption below it


def render_body(document, md_text, base_dir):
    names = style_names(document)
    table_style = table_style_name(document)
    width = text_width(document)
    in_landscape = False

    def table_after(k):
        """Index of a pipe table starting at the next non-blank line at or after k, else None."""
        while k < n and not lines[k].strip():
            k += 1
        if k + 1 < n and lines[k].strip().startswith("|") and is_table_separator(lines[k + 1]):
            return k
        return None
    bib_style = "Bibliography" if "Bibliography" in names else "Normal"

    lines = md_text.replace("\r\n", "\n").split("\n")
    n = len(lines)
    i = 0
    title_used = False
    in_references = False
    ref_anchor = None

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2).strip()
            if level == 1 and not title_used:
                p = document.add_paragraph(style="Title")
                add_runs(p, heading_text)
                title_used = True
            else:
                style = heading_style_for(level, names)
                p = document.add_paragraph(style=style)
                add_runs(p, heading_text)
            in_references = heading_text.strip().lower() == "references"
            if in_references and _CITER:
                ref_anchor = p._p
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # skip closing fence
            p = document.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for k, code_line in enumerate(code_lines):
                if k > 0:
                    p.add_run().add_break()
                run = p.add_run(code_line if code_line else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(8)
            continue

        if stripped.startswith(">"):
            p = document.add_paragraph(style="Normal")
            add_runs(p, stripped.lstrip(">").strip(), force_italic=True)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < n and is_table_separator(lines[i + 1]):
            header, rows, j = parse_table(lines, i)
            portrait_pt, landscape_pt = page_widths(document)
            landscape, size = plan_table(header, rows, portrait_pt, landscape_pt)
            if landscape and not in_landscape:
                open_section(document, True)
                in_landscape = True
            add_table(document, header, rows, table_style, size, landscape_pt if in_landscape else portrait_pt)
            spacer = document.add_paragraph(style="Normal")   # air between the table and what follows
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(10)
            spacer._p.get_or_add_pPr().append(OxmlElement("w:suppressLineNumbers"))
            if in_landscape:
                open_section(document, False)
                in_landscape = False
            i = j
            continue

        m = IMAGE_RE.match(stripped)
        if m:
            add_image(document, m.group(1), m.group(2), base_dir, width)
            i += 1
            continue

        if CAPTION_RE.match(stripped):
            t = table_after(i + 1)
            if t is not None and not in_landscape:
                # the caption stays with its table: open the landscape section before the caption
                header, rows, _ = parse_table(lines, t)
                portrait_pt, landscape_pt = page_widths(document)
                if plan_table(header, rows, portrait_pt, landscape_pt)[0]:
                    open_section(document, True)
                    in_landscape = True
            caption_text = TRAILING_ASIDE_RE.sub("", stripped)
            p = document.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = t is not None   # a table caption never ends a page alone
            add_runs(p, caption_text)
            i += 1
            continue

        if in_references and _CITER:
            # the ZOTERO_BIBL field replaces any hand-written reference list
            i += 1
            continue

        if in_references:
            p = document.add_paragraph(style=bib_style)
            add_runs(p, stripped)
            i += 1
            continue

        p = document.add_paragraph(style="Normal")
        add_runs(p, stripped)
        i += 1

    if _CITER and _CITER.cited:
        if ref_anchor is None:
            ref_anchor = document.add_paragraph(style=heading_style_for(1, names))._p
            add_runs(document.paragraphs[-1], "References")
        for para in reversed(_CITER.bibliography_paragraphs(bib_style)):
            ref_anchor.addnext(para)


def clear_body(document):
    body = document.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def strip_annotations(md_text):
    """The .docx is final text: drop ARTi-annotate CriticMarkup (comments keep their current span, suggestions
    keep the old text), so no `{==...==}{>>#aN ...<<}` leaks into the manuscript. The source .md is untouched."""
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "arti-annotate"))
    try:
        import annot
    except ImportError:
        return md_text
    finally:
        sys.path.pop(0)
    found = annot.parse(md_text)
    if found:
        open_n = sum(1 for a in found if a["status"] == "open")
        print("arti-docx: left out %d annotation(s) from the .docx (%d still open); the .md keeps them" % (len(found), open_n))
    return annot.strip(md_text)


def render(md_path, out_path, template_path, cite="none", project=None, cite_style=None):
    global _CITER
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()
    md_text = strip_annotations(md_text)

    _CITER = None
    if cite == "zotero":
        keys = zotero_fields.scan_keys(md_text)
        _CITER = zotero_fields.ZoteroCiter(zotero_fields.load_items(project, keys))

    document = docx.Document(template_path)
    clear_body(document)

    base_dir = os.path.dirname(os.path.abspath(md_path))
    render_body(document, md_text, base_dir)
    if _CITER and _CITER.cited:
        zotero_fields.set_style_pref(document, cite_style or zotero_fields.DEFAULT_STYLE)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    document.save(out_path)
    print("arti-docx: %s" % out_path)


def main():
    parser = argparse.ArgumentParser(description="Render a Markdown manuscript to a styled .docx")
    parser.add_argument("--md", required=True, help="path to the source Markdown file")
    parser.add_argument("--out", required=True, help="path to write the .docx to")
    parser.add_argument("--template", default=None, help="style template .docx (default: ARTi's academic template; the Zotero-carrying one with --cite zotero)")
    parser.add_argument("--cite", choices=["none", "zotero"], default="none", help="zotero: emit [LIT: key] markers as Zotero Word fields + a ZOTERO_BIBL reference list (needs enriched arti-lit.db)")
    parser.add_argument("--project", default=None, help="paper project root holding literature/arti-lit.db (default: cwd); used by --cite zotero")
    parser.add_argument("--cite-style", default=None, help="Zotero style id stored in the .docx (default: apa)")
    args = parser.parse_args()
    template = args.template or (ZOTERO_TEMPLATE if args.cite == "zotero" else DEFAULT_TEMPLATE)
    try:
        render(args.md, args.out, template, args.cite, args.project, args.cite_style)
    except zotero_fields.CiteError as e:
        raise SystemExit("arti-docx: %s" % e)


if __name__ == "__main__":
    main()
